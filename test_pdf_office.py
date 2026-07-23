"""Tests for PDF -> Office conversion (src/output/pdf_office.py).

A synthetic PDF is built with PyMuPDF: a title paragraph, a ruled 3-column
table (header + 2 data rows), and a trailing paragraph. Every expected value is
known, so the assertions lock in that:
  * tables come out as real rows x columns (XLSX cells, CSV rows, DOCX table);
  * prose survives as text (not swallowed by the table, not duplicated);
  * a scan-like PDF (no text layer) reports "no content" instead of a blank file.

No pytest in the venv — plain runner, same style as the other test_*.py files.
"""

import io

import fitz  # PyMuPDF

from src.output import pdf_office

W, H = 595, 842  # A4
COLS = [60, 240, 400, 520]      # 3 columns
ROWS = [140, 170, 200, 230]     # header + 2 data rows


def _build_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=W, height=H)
    page.insert_text((60, 90), "BAO CAO VAT TU", fontsize=16)  # title paragraph
    # Ruled table so find_tables keys on the grid.
    for x in COLS:
        page.draw_line(fitz.Point(x, ROWS[0]), fitz.Point(x, ROWS[-1]), width=0.8)
    for y in ROWS:
        page.draw_line(fitz.Point(COLS[0], y), fitz.Point(COLS[-1], y), width=0.8)
    cells = [
        ["Mo ta", "SL", "Thanh tien"],
        ["Thep D16", "10", "2500000"],
        ["Be tong C30", "5", "3750000"],
    ]
    for ri, row in enumerate(cells):
        for ci, val in enumerate(row):
            page.insert_text((COLS[ci] + 6, ROWS[ri] + 20), val, fontsize=10)
    page.insert_text((60, 300), "Ghi chu: nghiem thu tai hien truong.", fontsize=11)
    data = doc.tobytes(deflate=True)
    doc.close()
    return data


def _extract():
    pdf = _build_pdf()
    doc = fitz.open(stream=pdf, filetype="pdf")
    try:
        pages = pdf_office.extract_pages(doc, [0])
    finally:
        doc.close()
    return pdf, pages


def test_extract_finds_table_and_prose():
    _pdf, pages = _extract()
    assert len(pages) == 1
    items = pages[0]
    kinds = [it["kind"] for it in items]
    assert "table" in kinds, "the ruled table was not detected"
    assert "text" in kinds, "prose paragraphs were dropped"

    table = next(it for it in items if it["kind"] == "table")
    flat = "\n".join("|".join(r) for r in table["rows"])
    assert "Thep D16" in flat and "2500000" in flat, f"table cells missing: {flat!r}"
    # A row is rectangular (all rows same width).
    widths = {len(r) for r in table["rows"]}
    assert len(widths) == 1, f"table rows not rectangular: {widths}"

    prose = " ".join(it["text"] for it in items if it["kind"] == "text")
    assert "BAO CAO" in prose, "title paragraph missing"
    assert "nghiem thu" in prose, "trailing paragraph missing"
    # The table's cell text must NOT leak into the prose blocks.
    assert "Thanh tien" not in prose, "table header leaked into prose (double-counted)"


def test_csv_has_table_rows():
    pdf, _ = _extract()
    data, ok = pdf_office.convert(pdf, "csv")
    assert ok
    text = data.decode("utf-8-sig")
    assert "Thep D16" in text and "Be tong C30" in text
    # A data row keeps its three columns on one CSV line.
    assert any(line.count(",") >= 2 and "Thep D16" in line for line in text.splitlines()), text


def test_xlsx_has_table_cells():
    import openpyxl
    pdf, _ = _extract()
    data, ok = pdf_office.convert(pdf, "xlsx")
    assert ok
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    seen = {(str(c.value).strip() if c.value is not None else "") for row in ws.iter_rows() for c in row}
    assert "Thep D16" in seen and "2500000" in seen and "SL" in seen
    # The three columns of a data row occupy three separate cells in one row.
    hit = False
    for row in ws.iter_rows(values_only=True):
        vals = [("" if v is None else str(v).strip()) for v in row]
        if "Thep D16" in vals and "10" in vals and "2500000" in vals:
            hit = True
            break
    assert hit, "table row did not land as separate columns"


def test_docx_has_table_and_paragraph():
    from docx import Document
    pdf, _ = _extract()
    data, ok = pdf_office.convert(pdf, "docx")
    assert ok
    d = Document(io.BytesIO(data))
    assert d.tables, "no Word table emitted"
    cells = {c.text.strip() for t in d.tables for row in t.rows for c in row.cells}
    assert "Thep D16" in cells and "2500000" in cells
    paras = " ".join(p.text for p in d.paragraphs)
    assert "BAO CAO" in paras and "nghiem thu" in paras


def test_scan_pdf_reports_no_content():
    """A page with no text layer yields no content (caller shows 'is a scan')."""
    doc = fitz.open()
    doc.new_page(width=W, height=H)  # blank, no text
    pdf = doc.tobytes()
    doc.close()
    data, ok = pdf_office.convert(pdf, "xlsx")
    assert not ok and data == b""


def test_unknown_format_raises():
    pdf, _ = _extract()
    try:
        pdf_office.convert(pdf, "pptx")
        assert False, "expected ValueError for unsupported format"
    except ValueError:
        pass


if __name__ == "__main__":
    for fn in (
        test_extract_finds_table_and_prose,
        test_csv_has_table_rows,
        test_xlsx_has_table_cells,
        test_docx_has_table_and_paragraph,
        test_scan_pdf_reports_no_content,
        test_unknown_format_raises,
    ):
        fn()
        print(f"PASS {fn.__name__}")
    print("All pdf-office tests passed.")
